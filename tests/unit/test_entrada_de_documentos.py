"""Leitura dos documentos enviados e do formulário (sem chamar a API)."""

import io

import docx
import pytest

from app.api_comum import ErroParaUsuario, _nome_do_arquivo, ler_formulario, ler_json
from app.document_reader import ExtractionError, extrair_texto
from app.multipart_parser import extrair_boundary, parse_multipart


class TestExtrairTexto:
    def test_txt_em_utf8(self):
        assert extrair_texto("peticao.txt", "Pensão alimentícia".encode("utf-8")) == "Pensão alimentícia"

    def test_txt_do_bloco_de_notas_em_cp1252(self):
        assert extrair_texto("peticao.txt", "Pensão alimentícia".encode("cp1252")) == "Pensão alimentícia"

    def test_txt_com_bom_utf8(self):
        assert extrair_texto("peticao.txt", "\ufeffGuarda".encode("utf-8")) == "Guarda"

    def test_docx_inclui_paragrafos_e_tabelas(self):
        documento = docx.Document()
        documento.add_paragraph("Acordo de guarda compartilhada")
        tabela = documento.add_table(rows=1, cols=2)
        tabela.rows[0].cells[0].text = "Pai"
        tabela.rows[0].cells[1].text = "Fins de semana alternados"
        buffer = io.BytesIO()
        documento.save(buffer)

        texto = extrair_texto("acordo.docx", buffer.getvalue())
        assert "Acordo de guarda compartilhada" in texto
        assert "Pai | Fins de semana alternados" in texto

    def test_extensao_em_maiusculas_e_aceita(self):
        assert extrair_texto("PETICAO.TXT", b"texto") == "texto"

    @pytest.mark.parametrize("nome", ["planilha.xlsx", "foto.jpg", "sem_extensao"])
    def test_tipo_nao_aceito_tem_mensagem_clara(self, nome):
        with pytest.raises(ExtractionError, match="não é de um tipo aceito"):
            extrair_texto(nome, b"conteudo")

    def test_arquivo_vazio(self):
        with pytest.raises(ExtractionError, match="está vazio"):
            extrair_texto("peticao.txt", b"")

    def test_arquivo_so_com_espacos(self):
        with pytest.raises(ExtractionError, match="Não encontramos texto"):
            extrair_texto("peticao.txt", b"   \n\t  ")

    def test_pdf_corrompido(self):
        with pytest.raises(ExtractionError, match="Não conseguimos abrir o PDF"):
            extrair_texto("peticao.pdf", b"isto nao e um pdf")


def _corpo_multipart(boundary: str, partes: list[tuple[str, str | None, bytes]]) -> bytes:
    corpo = b""
    for nome, arquivo, conteudo in partes:
        disposicao = f'form-data; name="{nome}"' + (f'; filename="{arquivo}"' if arquivo else "")
        corpo += f"--{boundary}\r\nContent-Disposition: {disposicao}\r\n\r\n".encode() + conteudo + b"\r\n"
    return corpo + f"--{boundary}--\r\n".encode()


class TestMultipart:
    def test_extrai_boundary_com_e_sem_aspas(self):
        assert extrair_boundary("multipart/form-data; boundary=abc123") == b"abc123"
        assert extrair_boundary('multipart/form-data; boundary="abc123"') == b"abc123"
        assert extrair_boundary("application/json") is None

    def test_separa_campos_e_arquivos(self):
        corpo = _corpo_multipart("XYZ", [("tese", None, b"Partilha"), ("documentos", "a.txt", b"conteudo")])
        partes = parse_multipart(corpo, b"XYZ")
        assert [(p.name, p.filename, p.content) for p in partes] == [
            ("tese", None, b"Partilha"),
            ("documentos", "a.txt", b"conteudo"),
        ]


class TestFormulario:
    CT = "multipart/form-data; boundary=XYZ"

    def test_formulario_valido(self):
        corpo = _corpo_multipart("XYZ", [
            ("tese", None, "Não houve abandono".encode()),
            ("documentos", "peticao.txt", "Texto da petição".encode()),
        ])
        documentos, tese = ler_formulario(self.CT, corpo)
        assert tese == "Não houve abandono"
        assert documentos == {"peticao.txt": "Texto da petição"}

    def test_descarta_caminho_enviado_pelo_navegador(self):
        assert _nome_do_arquivo("C:\\Users\\cliente\\peticao.txt") == "peticao.txt"
        assert _nome_do_arquivo("/home/cliente/peticao.txt") == "peticao.txt"

    def test_sem_documentos(self):
        corpo = _corpo_multipart("XYZ", [("tese", None, b"Tese")])
        with pytest.raises(ErroParaUsuario) as erro:
            ler_formulario(self.CT, corpo)
        assert erro.value.status == 400

    def test_sem_tese(self):
        corpo = _corpo_multipart("XYZ", [("documentos", "a.txt", b"texto")])
        with pytest.raises(ErroParaUsuario) as erro:
            ler_formulario(self.CT, corpo)
        assert erro.value.status == 400

    def test_documento_invalido_explica_o_motivo(self):
        corpo = _corpo_multipart("XYZ", [("tese", None, b"Tese"), ("documentos", "foto.jpg", b"\x89PNG")])
        with pytest.raises(ErroParaUsuario, match="não é de um tipo aceito"):
            ler_formulario(self.CT, corpo)

    def test_envio_sem_boundary(self):
        with pytest.raises(ErroParaUsuario) as erro:
            ler_formulario("application/json", b"{}")
        assert erro.value.status == 400


class TestJson:
    def test_json_valido(self):
        assert ler_json('{"pergunta": "Por quê?"}'.encode()) == {"pergunta": "Por quê?"}

    @pytest.mark.parametrize("corpo", [b"nao e json", b"[1, 2]", b"\xff\xfe"])
    def test_json_invalido(self, corpo):
        with pytest.raises(ErroParaUsuario) as erro:
            ler_json(corpo)
        assert erro.value.status == 400
