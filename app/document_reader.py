"""Extração de texto dos documentos enviados pelo advogado.

Formatos aceitos: PDF (pypdf), Word .docx (python-docx) e texto (.txt/.md). As mensagens
de erro são mostradas diretamente na tela para o usuário, então falam de arquivos e do
que fazer — nunca de bibliotecas, codificação ou detalhes técnicos.
"""

from __future__ import annotations

import io

FORMATOS_ACEITOS = ("pdf", "docx", "txt", "md")


class ExtractionError(ValueError):
    """Levantado quando não foi possível extrair texto de um documento enviado."""


def extrair_texto(nome_arquivo: str, conteudo: bytes) -> str:
    extensao = nome_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_arquivo else ""

    if not conteudo:
        raise ExtractionError(f"O arquivo “{nome_arquivo}” está vazio.")

    if extensao in ("txt", "md"):
        texto = _ler_texto(nome_arquivo, conteudo)
    elif extensao == "pdf":
        texto = _ler_pdf(nome_arquivo, conteudo)
    elif extensao == "docx":
        texto = _ler_docx(nome_arquivo, conteudo)
    else:
        raise ExtractionError(
            f"O arquivo “{nome_arquivo}” não é de um tipo aceito. "
            "Envie documentos em PDF, Word (.docx) ou texto (.txt)."
        )

    if not texto.strip():
        raise ExtractionError(f"Não encontramos texto no arquivo “{nome_arquivo}”.")
    return texto


def _ler_texto(nome_arquivo: str, conteudo: bytes) -> str:
    # Arquivos .txt salvos no Bloco de Notas do Windows costumam vir em cp1252, não UTF-8.
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ExtractionError(
        f"Não conseguimos ler o texto de “{nome_arquivo}”. "
        "Salve o arquivo novamente como texto (.txt) e tente de novo."
    )


def _ler_pdf(nome_arquivo: str, conteudo: bytes) -> str:
    try:
        import pypdf
    except ImportError as exc:
        raise ExtractionError(
            f"Não conseguimos ler o PDF “{nome_arquivo}” neste momento. "
            "Envie o documento em Word (.docx) ou texto (.txt)."
        ) from exc

    try:
        leitor = pypdf.PdfReader(io.BytesIO(conteudo))
        if leitor.is_encrypted:
            leitor.decrypt("")
        texto = "\n".join(pagina.extract_text() or "" for pagina in leitor.pages)
    except Exception as exc:
        raise ExtractionError(
            f"Não conseguimos abrir o PDF “{nome_arquivo}”. Verifique se o arquivo não "
            "está corrompido ou protegido por senha."
        ) from exc

    if not texto.strip():
        raise ExtractionError(
            f"O PDF “{nome_arquivo}” parece ser uma digitalização (imagem), sem texto "
            "selecionável. Envie uma versão com texto ou copie o conteúdo para um arquivo .txt."
        )
    return texto


def _ler_docx(nome_arquivo: str, conteudo: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError(
            f"Não conseguimos ler o arquivo Word “{nome_arquivo}” neste momento. "
            "Envie o documento em PDF ou texto (.txt)."
        ) from exc

    try:
        documento = docx.Document(io.BytesIO(conteudo))
        partes = [paragrafo.text for paragrafo in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                partes.append(" | ".join(celula.text for celula in linha.cells))
    except Exception as exc:
        raise ExtractionError(
            f"Não conseguimos abrir o arquivo Word “{nome_arquivo}”. Verifique se ele não "
            "está corrompido. Arquivos .doc antigos precisam ser salvos como .docx."
        ) from exc

    return "\n".join(parte for parte in partes if parte.strip())
